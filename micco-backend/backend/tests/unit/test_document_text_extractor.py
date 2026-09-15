"""Unit tests for the lightweight (non-NexusRAG) document text extractor.

This is a separate, cheaper code path from app.services.document_parser
(docling/marker) — no LLM calls, plain text only — used by the n8n agent
email-reply endpoints. Everything here runs against real temp files;
nothing external to mock per .claude/rules/testing.md.
"""
from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document as DocxDocument

from app.services.document_text_extractor import (
    AGENT_CONTENT_MAX_CHARS,
    extract_full_text,
)


async def test_extract_full_text_reads_txt_file(tmp_path: Path):
    file_path = tmp_path / "note.txt"
    file_path.write_text("Nội dung ghi chú test.", encoding="utf-8")

    result = await extract_full_text(file_path, "txt")

    assert result.supported is True
    assert result.content == "Nội dung ghi chú test."
    assert result.truncated is False


async def test_extract_full_text_reads_md_file(tmp_path: Path):
    file_path = tmp_path / "note.md"
    file_path.write_text("# Tiêu đề\n\nNội dung.", encoding="utf-8")

    result = await extract_full_text(file_path, "md")

    assert result.supported is True
    assert "Tiêu đề" in result.content


async def test_extract_full_text_reads_docx_paragraphs(tmp_path: Path):
    file_path = tmp_path / "bao_cao.docx"
    doc = DocxDocument()
    doc.add_paragraph("Đoạn một.")
    doc.add_paragraph("Đoạn hai.")
    doc.save(str(file_path))

    result = await extract_full_text(file_path, "docx")

    assert result.supported is True
    assert result.content == "Đoạn một.\nĐoạn hai."


async def test_extract_full_text_reads_pdf_pages(tmp_path: Path):
    file_path = tmp_path / "bao_cao.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Noi dung PDF test.")
    pdf.save(str(file_path))
    pdf.close()

    result = await extract_full_text(file_path, "pdf")

    assert result.supported is True
    assert "Noi dung PDF test." in result.content


async def test_extract_full_text_reports_unsupported_type(tmp_path: Path):
    file_path = tmp_path / "sheet.xlsx"
    file_path.write_bytes(b"not a real xlsx")

    result = await extract_full_text(file_path, "xlsx")

    assert result.supported is False
    assert result.content is None


async def test_extract_full_text_truncates_content_over_the_cap(tmp_path: Path):
    file_path = tmp_path / "long.txt"
    file_path.write_text("a" * (AGENT_CONTENT_MAX_CHARS + 500), encoding="utf-8")

    result = await extract_full_text(file_path, "txt")

    assert result.truncated is True
    assert len(result.content) == AGENT_CONTENT_MAX_CHARS


async def test_extract_full_text_degrades_gracefully_for_corrupt_pdf(tmp_path: Path):
    """agent-content serves PENDING documents that haven't been through the
    main NexusRAG parsing pipeline yet, so a corrupt/truncated PDF is a real
    scenario — it must not become an unhandled 500."""
    file_path = tmp_path / "corrupt.pdf"
    file_path.write_bytes(b"not a real pdf, just garbage bytes 1234567890")

    result = await extract_full_text(file_path, "pdf")

    assert result.supported is False
    assert result.content is None
    assert result.truncated is False


async def test_extract_full_text_degrades_gracefully_for_corrupt_docx(tmp_path: Path):
    file_path = tmp_path / "corrupt.docx"
    file_path.write_bytes(b"not a docx, garbage bytes")

    result = await extract_full_text(file_path, "docx")

    assert result.supported is False
    assert result.content is None
    assert result.truncated is False
